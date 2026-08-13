#!/usr/bin/env python3
"""주간보고 엑셀에서 담당자별 활동 내역을 뽑아 CSV로 반영한다.

사용법:
    python3 scripts/extract_weekly_activity.py <담당자명> <엑셀경로> <weekOf>

예시:
    python3 scripts/extract_weekly_activity.py 권영제 ~/주간보고_0810.xlsx 2026-08-10

결과는 data/weekly_activity_<담당자명>.csv 에 저장된다.
같은 weekOf 로 다시 실행하면 해당 주차 행만 새 내용으로 교체되고,
다른 주차 기록은 그대로 남는다.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import re
import sys
import unicodedata
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - 실행 환경 안내용
    sys.exit("openpyxl 이 필요합니다: pip install openpyxl")

REPO_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = REPO_ROOT / "data"

# 출력 CSV 의 컬럼 순서
FIELDNAMES = ["week_of", "person", "date", "category", "farm", "activity", "plan", "note"]

# 엑셀 헤더에서 찾아볼 이름들. 앞쪽에 있을수록 우선 매칭된다.
HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "date": ("날짜", "일자", "일시", "수행일", "활동일", "방문일", "date"),
    "person": ("담당자", "담당", "성명", "이름", "작성자", "사원명", "person", "name"),
    "category": ("구분", "분류", "유형", "카테고리", "업무구분", "활동구분", "category"),
    "farm": ("농장", "거래처", "고객", "고객사", "방문처", "업체", "농가", "farm"),
    "activity": (
        "활동내용",
        "업무내용",
        "추진내용",
        "주요활동",
        "주요업무",
        "실적",
        "금주실적",
        "금주활동",
        "내용",
        "activity",
    ),
    "plan": ("차주계획", "익주계획", "향후계획", "다음주계획", "계획", "plan"),
    "note": ("비고", "특이사항", "메모", "note", "remark"),
}

# 헤더 후보를 찾을 때 훑어볼 최대 행 수
HEADER_SCAN_ROWS = 30


def normalize(value: object) -> str:
    """헤더 비교용으로 공백·기호를 걷어낸 문자열을 만든다."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    return re.sub(r"[\s()\[\]/·.\-_]+", "", text).lower()


def cell_text(value: object) -> str:
    """셀 값을 CSV 에 넣을 문자열로 바꾼다."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        if value.hour or value.minute:
            return value.strftime("%Y-%m-%d %H:%M")
        return value.strftime("%Y-%m-%d")
    if isinstance(value, dt.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def match_field(header: str) -> str | None:
    """헤더 한 칸이 어떤 필드에 해당하는지 판단한다."""
    key = normalize(header)
    if not key:
        return None
    # 완전 일치를 먼저 보고, 없으면 부분 일치로 넘어간다.
    for field, aliases in HEADER_ALIASES.items():
        if any(key == normalize(alias) for alias in aliases):
            return field
    for field, aliases in HEADER_ALIASES.items():
        if any(normalize(alias) in key for alias in aliases):
            return field
    return None


def find_header_row(rows: list[tuple]) -> tuple[int, dict[str, int]] | None:
    """헤더로 보이는 행과 필드 → 열 인덱스 매핑을 찾는다."""
    best: tuple[int, int, dict[str, int]] | None = None
    for row_idx, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        mapping: dict[str, int] = {}
        for col_idx, value in enumerate(row):
            field = match_field(value)
            if field and field not in mapping:
                mapping[field] = col_idx
        # 활동내용 계열이 없으면 주간보고 표로 보기 어렵다.
        if "activity" not in mapping or len(mapping) < 2:
            continue
        score = len(mapping)
        if best is None or score > best[1]:
            best = (row_idx, score, mapping)
    if best is None:
        return None
    return best[0], best[2]


def extract_sheet(sheet, person: str) -> list[dict[str, str]]:
    """시트 하나에서 담당자의 활동 행을 뽑아낸다."""
    rows = list(sheet.iter_rows(values_only=True))
    found = find_header_row(rows)
    if found is None:
        return []
    header_idx, mapping = found

    person_key = normalize(person)
    has_person_column = "person" in mapping
    records: list[dict[str, str]] = []
    last_date = ""

    for row in rows[header_idx + 1 :]:
        values = {field: cell_text(row[col]) if col < len(row) else "" for field, col in mapping.items()}

        if has_person_column:
            row_person = values.get("person", "")
            if row_person and person_key not in normalize(row_person):
                continue
            if not row_person and not records:
                # 담당자 열이 있는데 첫 행부터 비어 있으면 남의 표일 수 있으니 건너뛴다.
                continue

        activity = values.get("activity", "")
        if not activity:
            continue
        # 헤더가 반복되는 표가 있어 같은 문구가 다시 나오면 건너뛴다.
        if match_field(activity) == "activity":
            continue

        # 병합 셀 탓에 날짜가 비는 행은 위 행 날짜를 물려받는다.
        date = values.get("date", "") or last_date
        last_date = date or last_date

        records.append(
            {
                "date": date,
                "category": values.get("category", ""),
                "farm": values.get("farm", ""),
                "activity": activity,
                "plan": values.get("plan", ""),
                "note": values.get("note", ""),
            }
        )

    return records


def extract(path: Path, person: str, sheet_name: str | None) -> list[dict[str, str]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                sys.exit(f"시트를 찾을 수 없습니다: {sheet_name} (있는 시트: {', '.join(workbook.sheetnames)})")
            sheets = [workbook[sheet_name]]
        else:
            sheets = list(workbook.worksheets)

        records: list[dict[str, str]] = []
        for sheet in sheets:
            records.extend(extract_sheet(sheet, person))
        return records
    finally:
        workbook.close()


def dedupe(records: list[dict[str, str]]) -> list[dict[str, str]]:
    """여러 시트에 같은 내용이 중복으로 들어간 경우를 정리한다."""
    seen: set[tuple[str, ...]] = set()
    unique: list[dict[str, str]] = []
    for record in records:
        key = tuple(record[field] for field in ("date", "category", "farm", "activity"))
        if key in seen:
            continue
        seen.add(key)
        unique.append(record)
    return unique


def read_existing(csv_path: Path) -> list[dict[str, str]]:
    if not csv_path.exists():
        return []
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        return [{field: row.get(field, "") for field in FIELDNAMES} for row in csv.DictReader(handle)]


def write_csv(csv_path: Path, rows: list[dict[str, str]]) -> None:
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(rows)


def parse_week_of(value: str) -> str:
    try:
        return dt.date.fromisoformat(value).isoformat()
    except ValueError:
        sys.exit(f"weekOf 는 YYYY-MM-DD 형식이어야 합니다: {value}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="주간보고 엑셀에서 담당자 활동 내역을 뽑아 data/weekly_activity_<담당자명>.csv 에 반영한다.",
    )
    parser.add_argument("person", help="담당자명 (예: 권영제)")
    parser.add_argument("excel_path", help="주간보고 엑셀 파일 경로")
    parser.add_argument("week_of", help="주차 시작일 YYYY-MM-DD (예: 2026-08-10)")
    parser.add_argument("--sheet", help="특정 시트만 읽을 때 시트 이름", default=None)
    parser.add_argument("--out", help="출력 CSV 경로 (기본: data/weekly_activity_<담당자명>.csv)", default=None)
    parser.add_argument("--dry-run", action="store_true", help="CSV 를 쓰지 않고 뽑힌 내용만 보여준다")
    args = parser.parse_args(argv)

    excel_path = Path(args.excel_path).expanduser()
    if not excel_path.is_file():
        sys.exit(f"엑셀 파일을 찾을 수 없습니다: {excel_path}")

    week_of = parse_week_of(args.week_of)
    csv_path = Path(args.out).expanduser() if args.out else DATA_DIR / f"weekly_activity_{args.person}.csv"

    records = dedupe(extract(excel_path, args.person, args.sheet))
    if not records:
        sys.exit(
            f"'{args.person}' 의 활동 내역을 찾지 못했습니다. "
            "--sheet 로 시트를 지정하거나 엑셀의 헤더 이름을 확인해 주세요."
        )

    new_rows = [{"week_of": week_of, "person": args.person, **record} for record in records]

    if args.dry_run:
        print(f"[dry-run] {week_of} / {args.person} — {len(new_rows)}건")
        for row in new_rows:
            print("  " + " | ".join(row[field] for field in ("date", "category", "farm", "activity")))
        return 0

    # 같은 주차·담당자 행은 새로 뽑은 내용으로 교체하고, 나머지는 남긴다.
    kept = [row for row in read_existing(csv_path) if not (row["week_of"] == week_of and row["person"] == args.person)]
    merged = sorted(kept + new_rows, key=lambda row: (row["week_of"], row["person"], row["date"]))
    write_csv(csv_path, merged)

    shown = csv_path.relative_to(REPO_ROOT) if csv_path.is_relative_to(REPO_ROOT) else csv_path
    print(f"{shown}: {week_of} / {args.person} {len(new_rows)}건 반영 (전체 {len(merged)}건)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
